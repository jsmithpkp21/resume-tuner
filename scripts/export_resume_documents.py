#!/usr/bin/env python3
"""Generate DOCX and PDF resume artifacts from the assembled resume output.
This command runs the existing build pipeline once, then exports both formats in a
deterministic order, preferring processed default-template HTML when available and
falling back to the assembled markdown artifact otherwise.
"""

from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
import time
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

if __package__ in {None, ""}:
    import build_resume
    from _runtime_guard import (
        assert_not_blocked_runtime_input as _assert_not_blocked_runtime_input,
    )
    from select_skills import SKILLS_SEPARATOR
else:
    from scripts import build_resume
    from scripts._runtime_guard import (
        assert_not_blocked_runtime_input as _assert_not_blocked_runtime_input,
    )
    from scripts.select_skills import SKILLS_SEPARATOR


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--profile", type=Path, default=build_resume.DEFAULT_PROFILE)
    parser.add_argument(
        "--experience-db", type=Path, default=build_resume.DEFAULT_EXPERIENCE_DB
    )
    parser.add_argument(
        "--skills-matrix", type=Path, default=build_resume.DEFAULT_SKILLS_MATRIX
    )
    parser.add_argument("--job-url", type=str, default="")
    parser.add_argument("--job-text-file", type=Path, default=None)
    parser.add_argument("--target-role", type=str, default="")
    parser.add_argument(
        "--company",
        type=str,
        default="company",
        help="Company slug source for canonical output names (<company>_resume.*).",
    )
    parser.add_argument(
        "--output-dir", type=Path, default=build_resume.DEFAULT_OUTPUT_DIR
    )
    parser.add_argument(
        "--processing-mode",
        choices=("raw", "processed"),
        default="processed",
        help="Use processed mode by default for submission-ready output.",
    )
    parser.add_argument(
        "--template",
        type=str,
        default="modern",
        help="Secondary template pass-through for build pipeline consistency.",
    )
    parser.add_argument(
        "--pdf-filename",
        type=str,
        default=None,
        help="Optional PDF file name override under --output-dir.",
    )
    parser.add_argument(
        "--docx-filename",
        type=str,
        default=None,
        help="Optional DOCX file name override under --output-dir.",
    )
    parser.add_argument(
        "--allow-overflow-pdf",
        action="store_true",
        help=(
            "Keep the generated PDF even when it exceeds the default two-page "
            "submission guard. Useful for review/debug exports."
        ),
    )
    parser.add_argument(
        "--post-layout-cleanup",
        choices=("enabled", "disabled"),
        default="enabled",
        help=(
            "Apply deterministic post-layout cleanup before DOCX/PDF render. "
            "Set to 'disabled' for strict source fidelity/debug exports."
        ),
    )
    return parser.parse_args()


def _snake_case(value: str) -> str:
    token = re.sub(r"[^a-z0-9]+", "_", value.strip().lower())
    token = token.strip("_")
    return token or "company"


def _canonical_export_filename(company: str, extension: str) -> str:
    ext = extension.lower().lstrip(".")
    return f"{_snake_case(company)}_resume.{ext}"


def _staging_path(output_path: Path, *, label: str) -> Path:
    token = f"tmp-export-{os.getpid()}-{int(time.time() * 1000)}"
    return output_path.with_name(f".{output_path.name}.{token}.{label}")


def _remove_stale_legacy_exports(output_dir: Path, keep_paths: set[Path]) -> list[Path]:
    """Remove outdated legacy export aliases that can mask the current outputs."""
    removed: list[Path] = []
    for legacy_name in ("latest_resume_export.docx", "latest_resume_export.pdf"):
        legacy_path = output_dir / legacy_name
        if legacy_path in keep_paths or not legacy_path.exists():
            continue
        legacy_path.unlink()
        removed.append(legacy_path)
    return removed


def _require_python_docx() -> tuple[Any, Any, Any]:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is required for DOCX export. "
            f"Install dependencies from requirements.txt. Import error: {exc}"
        ) from exc
    return Document, Pt, Inches


def _require_reportlab() -> tuple[Any, Any, Any]:
    try:
        from reportlab.lib.pagesizes import LETTER
        from reportlab.pdfbase.pdfmetrics import stringWidth
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise RuntimeError(
            "reportlab is required for PDF export. "
            f"Install dependencies from requirements.txt. Import error: {exc}"
        ) from exc
    return LETTER, canvas, stringWidth


def _find_calibri_path(style: str = "Regular") -> Path | None:
    """Locate exact Calibri TTF via WSL Windows mount, then fontconfig."""
    file_map = {
        "Regular": "calibri.ttf",
        "Bold": "calibrib.ttf",
    }
    expected_file = file_map.get(style)
    if expected_file is None:
        return None

    # Prefer exact Windows font files to avoid style substitutions like Calibri Light.
    windows_path = Path("/mnt/c/Windows/Fonts") / expected_file
    if windows_path.exists():
        return windows_path

    try:
        result = subprocess.run(
            [
                "fc-list",
                f":family=Calibri:style={style}",
                "--format=%{file}\n",
            ],
            capture_output=True,
            text=True,
            timeout=5,
        )
        for line in result.stdout.splitlines():
            path = Path(line.strip())
            if path.exists() and path.name.lower() == expected_file:
                return path
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass
    return None


def _find_fontconfig_font_path(family: str, style: str) -> Path | None:
    """Locate a font file from fontconfig for a given family/style pair."""
    try:
        result = subprocess.run(
            [
                "fc-list",
                f":family={family}:style={style}",
                "--format=%{file}\n",
            ],
            capture_output=True,
            text=True,
            timeout=5,
        )
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return None

    for line in result.stdout.splitlines():
        path = Path(line.strip())
        if path.exists():
            return path
    return None


def _require_calibri_pdf_fonts() -> tuple[str, str]:
    """Register and return PDF font names, preferring Calibri then deterministic fallbacks."""
    strict_calibri = os.getenv("RESUME_PDF_STRICT_CALIBRI", "0") == "1"
    regular_path = _find_calibri_path("Regular")
    bold_path = _find_calibri_path("Bold")

    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "reportlab is required for PDF export. Install dependencies from requirements.txt"
        ) from exc

    regular_name = "Calibri"
    bold_name = "Calibri-Bold"
    registered = set(pdfmetrics.getRegisteredFontNames())

    if regular_path is not None and bold_path is not None:
        if regular_name not in registered:
            pdfmetrics.registerFont(TTFont(regular_name, str(regular_path)))
        if bold_name not in registered:
            pdfmetrics.registerFont(TTFont(bold_name, str(bold_path)))
        return regular_name, bold_name

    if strict_calibri:
        raise FileNotFoundError(
            "Strict Calibri mode is enabled but exact Calibri fonts were not found.\n"
            "Set RESUME_PDF_STRICT_CALIBRI=0 to allow fallback fonts, or install:\n"
            "  /mnt/c/Windows/Fonts/calibri.ttf\n"
            "  /mnt/c/Windows/Fonts/calibrib.ttf"
        )

    # CI/container fallback path keeps PDF export deterministic without proprietary fonts.
    fallback_candidates = (
        ("Carlito", "Carlito-Bold"),
        ("Liberation Sans", "Liberation Sans Bold"),
    )
    for family_regular, family_bold in fallback_candidates:
        fallback_regular = _find_fontconfig_font_path(family_regular, "Regular")
        fallback_bold = _find_fontconfig_font_path(family_regular, "Bold")
        if fallback_regular is None or fallback_bold is None:
            continue
        if family_regular not in registered:
            pdfmetrics.registerFont(TTFont(family_regular, str(fallback_regular)))
        if family_bold not in registered:
            pdfmetrics.registerFont(TTFont(family_bold, str(fallback_bold)))
        return family_regular, family_bold

    # Final fallback avoids hard dependency on fontconfig/fc-list in minimal containers.
    return "Helvetica", "Helvetica-Bold"


def _pipeline_markdown_path(output_dir: Path, processing_mode: str) -> Path:
    suffix = "processed" if processing_mode == "processed" else "raw"
    return output_dir / f"latest_resume_{suffix}.md"


def _pipeline_default_html_path(args: argparse.Namespace) -> Path:
    """Return the left-justified default-template HTML path from build output."""
    # Keep naming local to avoid coupling to build_resume private helpers.
    output_prefix = f"latest_resume_{args.processing_mode}"
    primary_template = "modern" if args.processing_mode == "processed" else "default"
    if primary_template == "default":
        return Path(args.output_dir) / f"{output_prefix}.html"

    secondary_template = args.template
    if secondary_template == primary_template:
        secondary_template = "default"
    if secondary_template == "default":
        secondary_prefix = f"latest_default_resume_{args.processing_mode}"
        return Path(args.output_dir) / f"{secondary_prefix}.html"

    # Fallback for custom template invocations: default to primary HTML.
    return Path(args.output_dir) / f"{output_prefix}.html"


def _run_build_pipeline(args: argparse.Namespace) -> Path:
    pipeline_args = argparse.Namespace(
        profile=args.profile,
        experience_db=args.experience_db,
        skills_matrix=args.skills_matrix,
        job_url=args.job_url,
        job_text_file=args.job_text_file,
        target_role=args.target_role,
        output_dir=args.output_dir,
        processing_mode=args.processing_mode,
        skip_markdown=False,
        template=args.template,
    )
    rc = build_resume.run_pipeline(pipeline_args)
    if rc != 0:
        raise RuntimeError(f"build pipeline failed with exit code {rc}")
    md_path = _pipeline_markdown_path(args.output_dir, args.processing_mode)
    if not md_path.exists():
        raise RuntimeError(f"expected markdown output not found: {md_path}")
    return md_path


# ---------------------------------------------------------------------------
# Shared markdown block parser
# ---------------------------------------------------------------------------
_SKIP_LINE_RE = re.compile(r"^Related skills:", re.IGNORECASE)
_BOLD_SPLIT_RE = re.compile(r"\*\*(.+?)\*\*")
_HIDDEN_HTML_CLASS_TOKENS = {"related-skills"}


class _ResumeHtmlBlockParser(HTMLParser):
    """Extract visible heading/paragraph/list blocks in DOM order."""

    # h2 sections whose <p> children should be treated as bullet-like items
    # so that post-layout cleanup rules apply uniformly across HTML and markdown.
    _BULLET_P_SECTIONS = {"key skills and expertise", "leadership & community"}

    def __init__(self) -> None:
        super().__init__()
        self.blocks: list[tuple[str, str]] = []
        self._capture_stack: list[dict[str, Any]] = []
        self._current_h2: str = ""

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        # Inline elements: inject bold markers / tab separator into current block
        if tag == "strong" and self._capture_stack:
            self._capture_stack[-1]["parts"].append("**")
            return
        if tag == "span" and self._capture_stack:
            self._capture_stack[-1]["parts"].append("	")
            return
        attrs_map = {name: (value or "") for name, value in attrs}
        class_tokens = [token for token in attrs_map.get("class", "").split() if token]
        if tag == "hr" and "header-divider" in class_tokens:
            self.blocks.append(("divider", ""))
            return
        if tag not in {"h1", "h2", "h3", "p", "li"}:
            return
        self._capture_stack.append({"tag": tag, "class": class_tokens, "parts": []})

    def handle_data(self, data: str) -> None:
        if not self._capture_stack:
            return
        self._capture_stack[-1]["parts"].append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "strong" and self._capture_stack:
            self._capture_stack[-1]["parts"].append("**")
            return
        if tag == "span":
            return  # span content already captured via handle_data
        if not self._capture_stack or self._capture_stack[-1]["tag"] != tag:
            return
        captured = self._capture_stack.pop()
        raw = "".join(captured["parts"])
        # Preserve \t separators injected from <span> tags; only normalize other whitespace.
        if "	" in raw:
            text = "	".join(
                " ".join(seg.split()) for seg in raw.split("	")
            ).strip()
        else:
            text = " ".join(raw.split()).strip()
        if not text:
            return

        class_tokens: list[str] = captured["class"]
        if any(token in _HIDDEN_HTML_CLASS_TOKENS for token in class_tokens):
            return
        if tag == "h1":
            self.blocks.append(("h1", text))
        elif tag == "h2":
            self._current_h2 = text.strip().lower()
            self.blocks.append(("h2", text))
        elif tag == "h3":
            self.blocks.append(("h3", text))
        elif tag == "li":
            self.blocks.append(("bullet", text))
        elif "resume-title" in class_tokens:
            self.blocks.append(("title", text))
        elif (
            "skills-category" in class_tokens
            or self._current_h2 in self._BULLET_P_SECTIONS
        ):
            # Treat skills rows and leadership/community paragraphs as bullets so
            # that _apply_post_layout_cleanup rules fire identically for HTML and
            # markdown sources.
            self.blocks.append(("bullet", text))
        else:
            self.blocks.append(("p", text))


def _iter_html_blocks(html_text: str) -> list[tuple[str, str]]:
    parser = _ResumeHtmlBlockParser()
    parser.feed(html_text)
    parser.close()
    return parser.blocks


def _iter_markdown_blocks(md_text: str) -> list[tuple[str, str]]:
    """Return (kind, text) pairs from markdown or HTML source text."""
    if "<html" in md_text.lower() or "<!doctype html" in md_text.lower():
        return _iter_html_blocks(md_text)

    blocks: list[tuple[str, str]] = []
    for raw in md_text.splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if _SKIP_LINE_RE.match(line):
            continue
        if line.startswith("# "):
            blocks.append(("h1", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
        elif line.startswith("### "):
            blocks.append(("h3", line[4:].strip()))
        elif re.match(r"^\s*---+\s*$", line):
            blocks.append(("divider", ""))
        elif line.startswith("[RESUME_TITLE] "):
            title_text = line[len("[RESUME_TITLE] ") :].strip()
            blocks.append(("title", _normalize_inline_text(title_text)))
        elif re.match(r"^\s*-\s+", line):
            blocks.append(("bullet", re.sub(r"^\s*-\s+", "", line, count=1).strip()))
        else:
            blocks.append(("p", line))
    return blocks


def _normalize_inline_text(text: str) -> str:
    return " ".join(text.replace("**", "").split())


def _has_single_word_wrap_tail(text: str, *, line_width: int) -> bool:
    normalized = _normalize_inline_text(text)
    wrapped = build_resume._summary_wrap_lines(normalized, line_width=line_width)
    return len(wrapped) > 1 and len(wrapped[-1].split()) == 1


def _drop_tail_skill_from_category_row(text: str) -> str:
    match = re.match(r"^\*\*(.+?):\*\*\s*(.+)$", text)
    if not match:
        return text
    category = match.group(1).strip()
    skills_text = match.group(2).strip()
    separator = SKILLS_SEPARATOR if SKILLS_SEPARATOR in skills_text else ", "
    if separator == SKILLS_SEPARATOR:
        skills_raw = [
            skill.strip()
            for skill in skills_text.split(SKILLS_SEPARATOR)
            if skill.strip()
        ]
    else:
        skills_raw = [
            skill.strip() for skill in re.split(r",\s*", skills_text) if skill.strip()
        ]
    if len(skills_raw) <= 1:
        return text
    return f"**{category}:** {separator.join(skills_raw[:-1])}"


def _trim_trailing_word(text: str) -> str:
    words = text.strip().split()
    if len(words) <= 6:
        return text
    shortened = " ".join(words[:-1]).rstrip(" ,;:")
    if text.strip().endswith((".", "!", "?")) and not shortened.endswith(
        (".", "!", "?")
    ):
        shortened += "."
    return shortened


def _job_context_has_community_signal(args: argparse.Namespace) -> bool:
    signal_terms = {
        "community",
        "dei",
        "diversity",
        "inclusion",
        "activism",
        "volunteer",
        "philanthropy",
        "outreach",
        "nonprofit",
    }
    source_parts = [args.target_role or "", args.company or ""]
    if args.job_text_file is not None:
        job_text_path = Path(args.job_text_file)
        _assert_not_blocked_runtime_input(job_text_path)
        if job_text_path.exists():
            source_parts.append(job_text_path.read_text(encoding="utf-8"))
    haystack = "\n".join(source_parts).lower()
    return any(term in haystack for term in signal_terms)


def _apply_post_layout_cleanup(
    source_text: str,
    *,
    args: argparse.Namespace,
) -> str:
    """Deterministically tighten assembled blocks before DOCX/PDF render."""
    blocks = _iter_markdown_blocks(source_text)
    if not blocks:
        return source_text

    has_mentoring_in_experience = False
    section = ""
    for kind, text in blocks:
        if kind == "h2":
            section = text.strip().lower()
            continue
        if section == "professional experience" and kind == "bullet":
            if "mentor" in _normalize_inline_text(text).lower():
                has_mentoring_in_experience = True
                break

    filtered_blocks: list[tuple[str, str]] = []
    section = ""
    removed_mentoring = False
    removed_philanthropy = False
    pending_leadership_detail_drop = False
    has_community_signal = _job_context_has_community_signal(args)
    philanthropy_terms = {
        "philanthropy",
        "volunteer",
        "community",
        "advocacy",
        "outreach",
        "nonprofit",
    }

    for kind, text in blocks:
        if kind == "h2":
            section = text.strip().lower()
            filtered_blocks.append((kind, text))
            continue

        if section == "key skills and expertise" and kind == "bullet":
            candidate = text
            if _has_single_word_wrap_tail(
                candidate, line_width=build_resume.DEFAULT_BULLET_LINE_WIDTH
            ):
                candidate = _drop_tail_skill_from_category_row(candidate)
            filtered_blocks.append((kind, candidate))
            continue

        if section == "professional experience" and kind == "bullet":
            candidate = text
            if _has_single_word_wrap_tail(
                candidate, line_width=build_resume.DEFAULT_BULLET_LINE_WIDTH
            ):
                candidate = _trim_trailing_word(candidate)
            filtered_blocks.append((kind, candidate))
            continue

        if section == "leadership & community" and kind == "bullet":
            if pending_leadership_detail_drop:
                pending_leadership_detail_drop = False
                continue
            lowered = _normalize_inline_text(text).lower()
            if (
                has_mentoring_in_experience
                and not removed_mentoring
                and "mentor" in lowered
            ):
                removed_mentoring = True
                pending_leadership_detail_drop = True
                continue
            if (
                not has_community_signal
                and not removed_philanthropy
                and any(term in lowered for term in philanthropy_terms)
            ):
                removed_philanthropy = True
                pending_leadership_detail_drop = True
                continue
            filtered_blocks.append((kind, text))
            continue

        filtered_blocks.append((kind, text))

    lines: list[str] = []
    for kind, text in filtered_blocks:
        if kind == "h1":
            lines.extend([f"# {text}", ""])
        elif kind == "h2":
            lines.extend([f"## {text}", ""])
        elif kind == "h3":
            lines.extend([f"### {text}", ""])
        elif kind == "title":
            lines.extend([f"[RESUME_TITLE] {text}", ""])
        elif kind == "divider":
            lines.extend(["---", ""])
        elif kind == "bullet":
            lines.append(f"- {text}")
        else:
            lines.append(text)

    return "\n".join(lines).strip() + "\n"


# ---------------------------------------------------------------------------
# DOCX renderer
# ---------------------------------------------------------------------------
def _add_rich_runs(
    paragraph: Any, text: str, base_size_pt: float, bold_base: bool = False
) -> None:
    """Add runs to a paragraph, honouring **bold** markers, stripping link syntax."""
    _, Pt, _ = _require_python_docx()
    text = _MD_LINK_RE.sub(r"\1", text)
    text = _MD_CODE_RE.sub(r"\1", text)
    parts = _BOLD_SPLIT_RE.split(text)
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        _set_docx_font_name(run, "Calibri")
        run.font.size = Pt(base_size_pt)
        # odd-indexed splits are the captured bold groups
        run.bold = bold_base or (i % 2 == 1)


def _set_docx_font_name(target: Any, font_name: str) -> None:
    """Force font name across Word script channels (ascii/hAnsi/eastAsia/cs)."""
    from docx.oxml.ns import qn

    target.font.name = font_name
    rpr = target._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), font_name)


def _render_docx(md_text: str, output_path: Path) -> None:
    Document, Pt, Inches = _require_python_docx()
    from docx.enum.text import WD_LINE_SPACING

    divider_after_pt = 12
    bullet_hanging_indent_pt = 14.4  # 0.20in

    doc = Document()
    # Narrow margins to match HTML template (0.5in each side)
    for section in doc.sections:
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
    # Base style
    normal = doc.styles["Normal"]
    _set_docx_font_name(normal, "Calibri")
    normal.font.size = Pt(11)

    def add_paragraph_bottom_border(
        paragraph: Any, *, color: str, size: int, space: int = 1
    ) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        p_pr = paragraph._p.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(size))
        bottom.set(qn("w:space"), str(space))
        bottom.set(qn("w:color"), color)
        p_bdr.append(bottom)

    def add_horizontal_rule(
        *, color: str, size: int, before_pt: float, after_pt: float
    ) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before_pt)
        p.paragraph_format.space_after = Pt(after_pt)
        p.add_run("")
        add_paragraph_bottom_border(p, color=color, size=size)

    def apply_word_body_paragraph_settings(paragraph: Any) -> None:
        """Match Word paragraph settings used in the approved sample output."""
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        # Keep body/list text at a fixed single-line equivalent to reduce drift.
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(11)

    previous_paragraph: Any | None = None
    seen_role_title = False

    for kind, text in _iter_markdown_blocks(md_text):
        if kind == "h1":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.bold = True
            _set_docx_font_name(run, "Calibri")
            run.font.size = Pt(18)
            previous_paragraph = p
        elif kind == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.bold = True
            _set_docx_font_name(run, "Calibri")
            run.font.size = Pt(12)
            add_paragraph_bottom_border(p, color="CCCCCC", size=8, space=0)
            previous_paragraph = p
        elif kind == "h3":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0 if not seen_role_title else 4)
            p.paragraph_format.space_after = Pt(0)
            if "	" in text:
                title_part, date_part = text.split("	", 1)
                title_plain = _strip_markdown_markup(title_part).strip()
                date_plain = _strip_markdown_markup(date_part).strip()
                run = p.add_run(title_plain)
                run.bold = True
                _set_docx_font_name(run, "Calibri")
                run.font.size = Pt(11)
                # Right tab stop at content width (7.5 in = 10800 twips)
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn

                pPr = p._p.get_or_add_pPr()
                tabs_el = OxmlElement("w:tabs")
                tab_el = OxmlElement("w:tab")
                tab_el.set(qn("w:val"), "right")
                tab_el.set(qn("w:pos"), "10800")
                tabs_el.append(tab_el)
                pPr.append(tabs_el)
                date_run = p.add_run("	" + date_plain)
                date_run.bold = False
                _set_docx_font_name(date_run, "Calibri")
                date_run.font.size = Pt(11)
            else:
                run = p.add_run(text)
                run.bold = True
                _set_docx_font_name(run, "Calibri")
                run.font.size = Pt(11)
            seen_role_title = True
            previous_paragraph = p
        elif kind == "title":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(_strip_markdown_markup(text))
            run.bold = True
            _set_docx_font_name(run, "Calibri")
            run.font.size = Pt(11)
            previous_paragraph = p
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            apply_word_body_paragraph_settings(p)
            # Force a deterministic hanging indent for wrapped bullet lines.
            p.paragraph_format.left_indent = Pt(bullet_hanging_indent_pt)
            p.paragraph_format.first_line_indent = Pt(-bullet_hanging_indent_pt)
            _add_rich_runs(p, text, 11)
            previous_paragraph = p
        elif kind == "divider":
            if previous_paragraph is None:
                add_horizontal_rule(
                    color="000000", size=12, before_pt=0, after_pt=divider_after_pt
                )
            else:
                add_paragraph_bottom_border(
                    previous_paragraph, color="000000", size=12, space=0
                )
                previous_paragraph.paragraph_format.space_after = Pt(divider_after_pt)
        else:  # p — contact line, date line, role summary, company line, skills
            p = doc.add_paragraph()
            apply_word_body_paragraph_settings(p)
            if "	" in text:
                # Company line: "**Company, Location**	Date Range"
                left_part, right_part = text.split("	", 1)
                left_plain = _strip_markdown_markup(left_part).strip()
                right_plain = _strip_markdown_markup(right_part).strip()
                run = p.add_run(left_plain)
                run.bold = True
                _set_docx_font_name(run, "Calibri")
                run.font.size = Pt(11)
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn

                pPr = p._p.get_or_add_pPr()
                tabs_el = OxmlElement("w:tabs")
                tab_el = OxmlElement("w:tab")
                tab_el.set(qn("w:val"), "right")
                tab_el.set(qn("w:pos"), "10800")
                tabs_el.append(tab_el)
                pPr.append(tabs_el)
                date_run = p.add_run("	" + right_plain)
                date_run.bold = True
                _set_docx_font_name(date_run, "Calibri")
                date_run.font.size = Pt(11)
            elif (skills_parts := _split_skills_category_line(text)) is not None:
                category, skills = skills_parts
                cat_run = p.add_run(f"{category} ")
                cat_run.bold = True
                _set_docx_font_name(cat_run, "Calibri")
                cat_run.font.size = Pt(11)
                skills_run = p.add_run(skills)
                skills_run.bold = False
                _set_docx_font_name(skills_run, "Calibri")
                skills_run.font.size = Pt(11)
            else:
                _add_rich_runs(p, text, 11)
            previous_paragraph = p
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


# ---------------------------------------------------------------------------
# PDF renderer
# ---------------------------------------------------------------------------
def _wrap_text_for_pdf(
    text: str, max_width: float, font_name: str, font_size: int
) -> list[str]:
    _LETTER, _canvas, string_width = _require_reportlab()
    words = text.split()
    if not words:
        return [""]
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = f"{current} {word}"
        if string_width(candidate, font_name, font_size) <= max_width:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def _wrap_skills_category_for_pdf(
    bold_prefix: str,
    regular_text: str,
    max_width: float,
    fn_bold: str,
    fn_regular: str,
    font_size: int,
) -> list[tuple[str, str]]:
    """Wrap a skills-category line using the correct font per segment.
    Returns a list of (bold_part, regular_part) tuples, one per output line.
    The bold_prefix is measured with fn_bold; remaining words with fn_regular,
    so wrap points reflect the actual rendered glyph widths.
    """
    _LETTER, _canvas, sw = _require_reportlab()
    prefix_w = sw(bold_prefix, fn_bold, font_size)
    words = regular_text.split()
    lines: list[tuple[str, str]] = []
    # --- First line: starts with bold prefix ---
    current_regular = ""
    current_w = prefix_w
    for i, word in enumerate(words):
        sep = " " if current_regular else ""
        word_w = sw(sep + word, fn_regular, font_size)
        if current_w + word_w <= max_width:
            current_regular += sep + word
            current_w += word_w
        else:
            lines.append((bold_prefix, current_regular))
            remaining = words[i:]
            break
    else:
        # All words fit on first line
        lines.append((bold_prefix, current_regular))
        return lines
    # --- Subsequent lines: all regular ---
    if remaining:
        current = remaining[0]
        for word in remaining[1:]:
            candidate = f"{current} {word}"
            if sw(candidate, fn_regular, font_size) <= max_width:
                current = candidate
            else:
                lines.append(("", current))
                current = word
        lines.append(("", current))
    return lines


_MD_LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]*\)")
_MD_CODE_RE = re.compile(r"`([^`]+)`")
_SKILLS_CATEGORY_RE = re.compile(r"^\*\*(.+?:)\*\*\s*(.+)$")


def _strip_markdown_markup(text: str) -> str:
    """Strip bold, links, and inline code markers for plain-text renderers."""
    text = _BOLD_SPLIT_RE.sub(r"\1", text)
    text = _MD_LINK_RE.sub(r"\1", text)
    text = _MD_CODE_RE.sub(r"\1", text)
    return text


def _split_skills_category_line(text: str) -> tuple[str, str] | None:
    """Return (category, skills) for lines like '**Category:** skill1 • skill2'."""
    match = _SKILLS_CATEGORY_RE.match(text.strip())
    if not match:
        return None
    return match.group(1).strip(), match.group(2).strip()


def _render_pdf(
    md_text: str, output_path: Path, *, enforce_page_limit: bool = True
) -> None:
    import io as _io

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.unlink(missing_ok=True)

    LETTER, canvas_mod, _sw = _require_reportlab()
    fn_regular, fn_bold = _require_calibri_pdf_fonts()
    buf = _io.BytesIO()
    pdf = canvas_mod.Canvas(buf, pagesize=LETTER)
    width, height = LETTER
    # Keep near-HTML width while preserving a stable two-page PDF envelope.
    margin_x: float = 36
    margin_top: float = 36
    margin_bottom: float = 36
    y: float = height - margin_top
    page_count = 1
    content_width = width - 2 * margin_x
    seen_divider = False
    last_text_baseline_y: float | None = None

    def new_page() -> None:
        nonlocal y, page_count
        pdf.showPage()
        page_count += 1
        y = height - margin_top

    def ensure_space(needed_pts: float) -> None:
        if y - needed_pts < margin_bottom:
            new_page()

    def ensure_line_space(line_height: float) -> None:
        ensure_space(line_height)

    def draw_text(
        x: float, y_pos: float, value: str, font_name: str, font_size: int
    ) -> None:
        pdf.setFont(font_name, font_size)
        pdf.drawString(x, y_pos, value)

    def draw_horizontal_rule(
        *, before_pts: float, after_pts: float, width: float = 1.0, gray: float = 0.0
    ) -> None:
        """Draw a rule in its own vertical row so it never overlaps text."""
        nonlocal y
        ensure_space(before_pts + 1 + after_pts)
        y -= before_pts
        pdf.setStrokeColorRGB(gray, gray, gray)
        pdf.setLineWidth(width)
        pdf.line(margin_x, y, margin_x + content_width, y)
        y -= after_pts

    for kind, text in _iter_markdown_blocks(md_text):
        plain = _strip_markdown_markup(text)
        if kind == "h1":
            fn, fs, lh = fn_bold, 18, 15.2
            ensure_space(lh + 4)
            draw_text(margin_x, y, plain, fn, fs)
            last_text_baseline_y = y
            y -= lh
        elif kind == "h2":
            fn, fs, lh = fn_bold, 12, 15
            ensure_space(lh + 16)
            y -= 5
            heading_y = y
            draw_text(margin_x, heading_y, plain, fn, fs)
            last_text_baseline_y = heading_y
            line_y = heading_y - 2
            pdf.setStrokeColorRGB(0.8, 0.8, 0.8)
            pdf.setLineWidth(0.8)
            pdf.line(margin_x, line_y, margin_x + content_width, line_y)
            y = line_y - 14
        elif kind == "h3":
            fn_bold_name, fn_reg_name, fs, lh = fn_bold, fn_regular, 11, 13.2
            ensure_space(lh + 2)
            y -= 1  # space before job title
            if "	" in text:
                title_part, date_part = text.split("	", 1)
                title_plain = _strip_markdown_markup(title_part).strip()
                date_plain = _strip_markdown_markup(date_part).strip()
            else:
                title_plain = plain
                date_plain = None
            draw_text(margin_x, y, title_plain, fn_bold_name, fs)
            last_text_baseline_y = y
            if date_plain:
                _sw = _require_reportlab()[2]
                dw = _sw(date_plain, fn_reg_name, 11)
                draw_text(margin_x + content_width - dw, y, date_plain, fn_reg_name, 11)
            y -= lh
        elif kind == "title":
            fn, fs, lh = fn_bold, 11, 13
            ensure_space(lh + 2)
            draw_text(margin_x, y, plain, fn, fs)
            last_text_baseline_y = y
            y -= lh
        elif kind == "divider":
            seen_divider = True
            # Match h2 behavior: line sits just under the preceding text and leaves 16pt after.
            if last_text_baseline_y is None:
                draw_horizontal_rule(before_pts=0, after_pts=16, width=1.0, gray=0.0)
            else:
                line_y = last_text_baseline_y - 2
                pdf.setStrokeColorRGB(0.0, 0.0, 0.0)
                pdf.setLineWidth(1.0)
                pdf.line(margin_x, line_y, margin_x + content_width, line_y)
                y = line_y - 16
        elif kind == "bullet":
            fn, fn_bold_name, fs, lh = fn_regular, fn_bold, 11, 12.0
            bullet_x = margin_x + 1
            text_x = margin_x + 14
            avail = content_width - (text_x - margin_x)
            lines: list[str]
            if (skills_parts := _split_skills_category_line(text)) is not None:
                category, skills = skills_parts
                category_with_space = f"{category} "
                _sw = _require_reportlab()[2]
                category_w = _sw(category_with_space, fn_bold_name, fs)
                skills_w = _sw(skills, fn, fs)
                if category_w + skills_w <= avail:
                    ensure_space(lh + 1)
                    ensure_line_space(lh)
                    draw_text(bullet_x, y, "•", fn, fs)
                    draw_text(text_x, y, category_with_space, fn_bold_name, fs)
                    draw_text(text_x + category_w, y, skills, fn, fs)
                    last_text_baseline_y = y
                    y -= lh
                else:
                    wrapped_lines = _wrap_skills_category_for_pdf(
                        category_with_space,
                        skills,
                        avail,
                        fn_bold_name,
                        fn,
                        fs,
                    )
                    ensure_space(len(wrapped_lines) * lh + 1)
                    for i, (bold_text, regular_text) in enumerate(wrapped_lines):
                        ensure_line_space(lh)
                        if i == 0:
                            draw_text(bullet_x, y, "•", fn, fs)
                        if bold_text:
                            draw_text(text_x, y, bold_text, fn_bold_name, fs)
                        if regular_text:
                            bold_w = (
                                _sw(bold_text, fn_bold_name, fs) if bold_text else 0
                            )
                            draw_text(text_x + bold_w, y, regular_text, fn, fs)
                        last_text_baseline_y = y
                        y -= lh
                y -= 0.2
            else:
                text_font = fn
                if "**" in text:
                    bold_parts = _BOLD_SPLIT_RE.split(text)
                    # Fully-bold markdown bullet line: **...**
                    if (
                        len(bold_parts) == 3
                        and not bold_parts[0].strip()
                        and not bold_parts[2].strip()
                    ):
                        text_font = fn_bold_name
                        bullet_text = _strip_markdown_markup(bold_parts[1])
                        lines = _wrap_text_for_pdf(bullet_text, avail, text_font, fs)
                    else:
                        lines = _wrap_text_for_pdf(plain, avail, text_font, fs)
                else:
                    lines = _wrap_text_for_pdf(plain, avail, text_font, fs)
                ensure_space(len(lines) * lh + 1)
                for i, line in enumerate(lines):
                    ensure_line_space(lh)
                    if i == 0:
                        draw_text(bullet_x, y, "•", text_font, fs)
                    draw_text(text_x, y, line, text_font, fs)
                    last_text_baseline_y = y
                    y -= lh
                y -= 0.2
        else:  # p — contact, date, role summary, company line, skills
            fn, fn_bold_name, fs, lh = fn_regular, fn_bold, 11, 11.0
            post_gap = 0 if not seen_divider else 0.2
            if "	" in text:
                # Company line: "**Company, Location**	Date Range"
                left_part, right_part = text.split("	", 1)
                left_plain = _strip_markdown_markup(left_part).strip()
                right_plain = _strip_markdown_markup(right_part).strip()
                ensure_space(lh + 1)
                _sw = _require_reportlab()[2]
                rw = _sw(right_plain, fn_bold_name, 11)
                draw_text(margin_x, y, left_plain, fn_bold_name, fs)
                last_text_baseline_y = y
                draw_text(
                    margin_x + content_width - rw, y, right_plain, fn_bold_name, 11
                )
                y -= lh
                y -= post_gap
            elif (skills_parts := _split_skills_category_line(text)) is not None:
                category, skills = skills_parts
                category_with_space = f"{category} "
                _sw = _require_reportlab()[2]
                category_w = _sw(category_with_space, fn_bold_name, fs)
                skills_w = _sw(skills, fn, fs)
                ensure_space(lh + 1)
                if category_w + skills_w <= content_width:
                    draw_text(margin_x, y, category_with_space, fn_bold_name, fs)
                    draw_text(margin_x + category_w, y, skills, fn, fs)
                    last_text_baseline_y = y
                    y -= lh
                else:
                    wrapped_lines = _wrap_skills_category_for_pdf(
                        category_with_space,
                        skills,
                        content_width,
                        fn_bold_name,
                        fn,
                        fs,
                    )
                    ensure_space(len(wrapped_lines) * lh + 1)
                    for bold_text, regular_text in wrapped_lines:
                        ensure_line_space(lh)
                        if bold_text:
                            draw_text(margin_x, y, bold_text, fn_bold_name, fs)
                        if regular_text:
                            bold_w = (
                                _sw(bold_text, fn_bold_name, fs) if bold_text else 0
                            )
                            draw_text(margin_x + bold_w, y, regular_text, fn, fs)
                        last_text_baseline_y = y
                        y -= lh
                y -= post_gap
            else:
                # Check for inline bold markers (e.g. degree line: **Bachelor of Science...**)
                if "**" in text:
                    _sw_fn = _require_reportlab()[2]
                    # Split on bold markers: odd-indexed parts are bold
                    bold_parts = _BOLD_SPLIT_RE.split(text)
                    # If entirely bold (single bold group wrapping all content), render bold
                    if (
                        len(bold_parts) == 3
                        and not bold_parts[0].strip()
                        and not bold_parts[2].strip()
                    ):
                        bold_text = _strip_markdown_markup(bold_parts[1])
                        lines = _wrap_text_for_pdf(
                            bold_text, content_width, fn_bold_name, fs
                        )
                        ensure_space(len(lines) * lh + 1)
                        for line in lines:
                            ensure_line_space(lh)
                            draw_text(margin_x, y, line, fn_bold_name, fs)
                            last_text_baseline_y = y
                            y -= lh
                    else:
                        # Mixed bold/regular: render on a single line (no wrapping for mixed)
                        ensure_space(lh + 1)
                        x_cursor = margin_x
                        for i, part in enumerate(bold_parts):
                            if not part:
                                continue
                            part_plain = _strip_markdown_markup(part)
                            is_bold = i % 2 == 1
                            cur_fn = fn_bold_name if is_bold else fn
                            pw = _sw_fn(part_plain, cur_fn, fs)
                            draw_text(x_cursor, y, part_plain, cur_fn, fs)
                            x_cursor += pw
                        last_text_baseline_y = y
                        y -= lh
                    y -= post_gap
                else:
                    lines = _wrap_text_for_pdf(plain, content_width, fn, fs)
                    ensure_space(len(lines) * lh + 1)
                    for line in lines:
                        ensure_line_space(lh)
                        draw_text(margin_x, y, line, fn, fs)
                        last_text_baseline_y = y
                        y -= lh
                    y -= post_gap
    pdf.save()
    if enforce_page_limit and page_count > 2:
        output_path.unlink(missing_ok=True)
        raise RuntimeError(
            f"PDF exceeded two-page limit ({page_count} pages): {output_path}"
        )
    output_path.write_bytes(buf.getvalue())


# ---------------------------------------------------------------------------
# Fragment quality guard
# ---------------------------------------------------------------------------
def _contains_trailing_connector_fragment(md_text: str) -> bool:
    bad_tail = {"and", "or", "to", "with", "for", "of", "in", "a", "an", "the"}
    for _kind, text in _iter_markdown_blocks(md_text):
        words = re.findall(r"[A-Za-z]+", text.lower())
        if words and words[-1] in bad_tail:
            return True
    return False


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
def run() -> int:
    args = parse_args()
    try:
        markdown_path = _run_build_pipeline(args)
        md_text = markdown_path.read_text(encoding="utf-8")
        render_source_text = md_text
        using_html_source = False
        html_source_path = _pipeline_default_html_path(args)
        if html_source_path.exists():
            render_source_text = html_source_path.read_text(encoding="utf-8")
            using_html_source = True
        post_layout_cleanup = getattr(args, "post_layout_cleanup", "enabled")
        if post_layout_cleanup == "enabled":
            render_source_text = _apply_post_layout_cleanup(
                render_source_text,
                args=args,
            )
        if _contains_trailing_connector_fragment(render_source_text):
            print(
                "WARNING: render source contains trailing connector fragments; review final DOCX/PDF for awkward wraps",
                file=sys.stderr,
            )
        docx_name = args.docx_filename or _canonical_export_filename(
            args.company, "docx"
        )
        pdf_name = args.pdf_filename or _canonical_export_filename(args.company, "pdf")
        docx_output = Path(args.output_dir) / docx_name
        pdf_output = Path(args.output_dir) / pdf_name
        if docx_output.resolve(strict=False) == pdf_output.resolve(strict=False):
            raise RuntimeError(
                "DOCX and PDF outputs must be different files; "
                "choose distinct --docx-filename and --pdf-filename values."
            )
        docx_output.parent.mkdir(parents=True, exist_ok=True)
        pdf_output.parent.mkdir(parents=True, exist_ok=True)

        staged_docx = _staging_path(docx_output, label="docx")
        staged_pdf = _staging_path(pdf_output, label="pdf")
        backup_docx = _staging_path(docx_output, label="docx.bak")
        backup_pdf = _staging_path(pdf_output, label="pdf.bak")
        docx_existed_before_finalize = docx_output.exists()
        pdf_existed_before_finalize = pdf_output.exists()

        try:
            _render_docx(render_source_text, staged_docx)
            _render_pdf(
                render_source_text,
                staged_pdf,
                enforce_page_limit=not args.allow_overflow_pdf,
            )

            if docx_existed_before_finalize:
                docx_output.replace(backup_docx)
            if pdf_existed_before_finalize:
                pdf_output.replace(backup_pdf)

            staged_docx.replace(docx_output)
            staged_pdf.replace(pdf_output)
        except Exception:
            staged_docx.unlink(missing_ok=True)
            staged_pdf.unlink(missing_ok=True)

            if backup_docx.exists():
                backup_docx.replace(docx_output)
            elif not docx_existed_before_finalize:
                docx_output.unlink(missing_ok=True)
            if backup_pdf.exists():
                backup_pdf.replace(pdf_output)
            elif not pdf_existed_before_finalize:
                pdf_output.unlink(missing_ok=True)
            raise
        # Backup cleanup is best-effort: a failure here must not undo a successful export.
        for backup in (backup_docx, backup_pdf):
            try:
                backup.unlink(missing_ok=True)
            except OSError as _e:
                print(
                    f"WARNING: could not remove backup file {backup}: {_e}",
                    file=sys.stderr,
                )
        removed_legacy_outputs = _remove_stale_legacy_exports(
            Path(args.output_dir), {docx_output, pdf_output}
        )
        print("Resume exports written:")
        print(f"  {docx_output}")
        print(f"  {pdf_output}")
        if removed_legacy_outputs:
            print("  removed stale legacy exports:")
            for removed_path in removed_legacy_outputs:
                print(f"    {removed_path}")
        print(f"  source markdown: {markdown_path}")
        if using_html_source:
            print(f"  source html: {html_source_path}")
        return 0
    except Exception as exc:  # noqa: BLE001
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(run())
