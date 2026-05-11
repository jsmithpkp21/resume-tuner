"""DOCX + PDF renderers for cover letters.

These are intentionally separate from the resume renderers in
``scripts/document_export.py``. Cover letters are paragraph-shaped:
no section dividers, no bullets, no right-aligned date stops, no
post-layout content trimming. Reusing the resume renderer would
produce a mis-styled letter (h2 borders, bullet glyphs, etc.).

Both renderers share font resolution, page-size helpers, and font
registration with the resume renderer via ``document_export``.
"""

from __future__ import annotations

import io as _io
import sys
from pathlib import Path
from typing import Any

from .cover_letter_template import CoverLetterIR
from .document_export import (
    require_calibri_pdf_fonts,
    require_python_docx,
    require_reportlab,
    set_docx_font_name,
    wrap_text_for_pdf,
)


def render_docx(ir: CoverLetterIR, output_path: Path) -> None:
    Document, Pt, Inches = require_python_docx()
    from docx.enum.text import WD_LINE_SPACING

    doc = Document()
    for section in doc.sections:
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
    normal = doc.styles["Normal"]
    set_docx_font_name(normal, "Calibri")
    normal.font.size = Pt(11)

    def add_paragraph(
        text: str, *, bold: bool = False, space_after_pt: float = 0.0
    ) -> Any:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after_pt)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        if text:
            run = p.add_run(text)
            run.bold = bold
            set_docx_font_name(run, "Calibri")
            run.font.size = Pt(11)
        return p

    def add_blank_line() -> None:
        add_paragraph("", space_after_pt=0)

    if ir.sender_name:
        add_paragraph(ir.sender_name, bold=True)
    for line in ir.sender_lines:
        if line:
            add_paragraph(line)
    add_blank_line()
    add_paragraph(ir.date_line)
    add_blank_line()
    for line in ir.recipient_lines:
        if line:
            add_paragraph(line)
    add_blank_line()
    add_paragraph(ir.greeting)
    add_blank_line()
    for paragraph in ir.body.all_paragraphs():
        text = paragraph.strip()
        if not text:
            continue
        add_paragraph(text)
        add_blank_line()
    add_paragraph(ir.closing)
    add_blank_line()
    if ir.sender_name:
        add_paragraph(ir.sender_name)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def render_pdf(
    ir: CoverLetterIR,
    output_path: Path,
    *,
    enforce_page_limit: bool = False,
) -> None:
    """Render the cover letter to PDF.

    ``enforce_page_limit=False`` (default) lets the renderer overflow
    onto additional pages and emits a ``WARNING:`` line to stderr when
    that happens (the page count is not surfaced to callers).
    Set to ``True`` to raise ``CoverLetterPageLimitError`` when the
    output exceeds one page (parity with resume's strict mode).
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.unlink(missing_ok=True)

    LETTER, canvas_mod, _ = require_reportlab()
    fn_regular, fn_bold = require_calibri_pdf_fonts()

    buf = _io.BytesIO()
    pdf = canvas_mod.Canvas(buf, pagesize=LETTER)
    width, height = LETTER

    margin_x = 72.0
    margin_top = 72.0
    margin_bottom = 72.0
    line_height = 14.0
    blank_line_height = 7.0
    body_paragraph_gap = 7.0

    content_width = width - 2 * margin_x
    y = height - margin_top
    page_count = 1

    def new_page() -> None:
        nonlocal y, page_count
        pdf.showPage()
        page_count += 1
        y = height - margin_top

    def ensure_space(needed: float) -> None:
        if y - needed < margin_bottom:
            new_page()

    def draw_line(text: str, *, font_name: str, font_size: int = 11) -> None:
        nonlocal y
        if not text:
            return
        wrapped = wrap_text_for_pdf(text, content_width, font_name, font_size)
        for segment in wrapped:
            ensure_space(line_height)
            pdf.setFont(font_name, font_size)
            pdf.drawString(margin_x, y, segment)
            y -= line_height

    def draw_blank() -> None:
        nonlocal y
        ensure_space(blank_line_height)
        y -= blank_line_height

    if ir.sender_name:
        draw_line(ir.sender_name, font_name=fn_bold)
    for line in ir.sender_lines:
        if line:
            draw_line(line, font_name=fn_regular)
    draw_blank()
    draw_line(ir.date_line, font_name=fn_regular)
    draw_blank()
    for line in ir.recipient_lines:
        if line:
            draw_line(line, font_name=fn_regular)
    draw_blank()
    draw_line(ir.greeting, font_name=fn_regular)
    draw_blank()
    for paragraph in ir.body.all_paragraphs():
        text = paragraph.strip()
        if not text:
            continue
        draw_line(text, font_name=fn_regular)
        ensure_space(body_paragraph_gap)
        y -= body_paragraph_gap
    draw_line(ir.closing, font_name=fn_regular)
    draw_blank()
    if ir.sender_name:
        draw_line(ir.sender_name, font_name=fn_regular)

    pdf.save()

    if page_count > 1:
        if enforce_page_limit:
            raise CoverLetterPageLimitError(
                f"Cover letter exceeded one page ({page_count} pages rendered)."
            )
        print(
            f"WARNING: cover letter rendered to {page_count} pages "
            f"(target: 1 page). Tighten the body or lower --word-budget.",
            file=sys.stderr,
        )

    output_path.write_bytes(buf.getvalue())


class CoverLetterPageLimitError(RuntimeError):
    """Raised when ``render_pdf(..., enforce_page_limit=True)`` overflows."""


def render_html(ir: CoverLetterIR) -> str:
    """Simple HTML for review-mode artifacts.

    Mirrors the markdown layout: stacked paragraphs, no headings, no
    bullets, sender name in bold.
    """
    parts: list[str] = [
        "<!doctype html>",
        '<html lang="en"><head><meta charset="utf-8">',
        f"<title>Cover Letter — {_escape(ir.sender_name)}</title>",
        "<style>",
        "body { font-family: Calibri, 'Segoe UI', Arial, sans-serif; "
        "font-size: 11pt; max-width: 7.5in; margin: 1in auto; line-height: 1.3; }",
        "p { margin: 0 0 0.5em 0; }",
        ".sender-name, .signature-name { font-weight: bold; }",
        ".body-paragraph { margin-bottom: 0.8em; }",
        "</style></head><body>",
    ]
    if ir.sender_name:
        parts.append(f'<p class="sender-name">{_escape(ir.sender_name)}</p>')
    for line in ir.sender_lines:
        if line:
            parts.append(f"<p>{_escape(line)}</p>")
    parts.append("<p>&nbsp;</p>")
    parts.append(f"<p>{_escape(ir.date_line)}</p>")
    parts.append("<p>&nbsp;</p>")
    for line in ir.recipient_lines:
        if line:
            parts.append(f"<p>{_escape(line)}</p>")
    parts.append("<p>&nbsp;</p>")
    parts.append(f"<p>{_escape(ir.greeting)}</p>")
    parts.append("<p>&nbsp;</p>")
    for paragraph in ir.body.all_paragraphs():
        text = paragraph.strip()
        if not text:
            continue
        parts.append(f'<p class="body-paragraph">{_escape(text)}</p>')
    parts.append(f"<p>{_escape(ir.closing)}</p>")
    parts.append("<p>&nbsp;</p>")
    if ir.sender_name:
        parts.append(f'<p class="signature-name">{_escape(ir.sender_name)}</p>')
    parts.append("</body></html>")
    return "\n".join(parts) + "\n"


def _escape(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )
